from datetime import datetime, timezone as zona_utc
from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import load_workbook

from apps.core.exportacion import TIPO_XLSX, generar_xlsx, respuesta_xlsx


def _abrir(contenido: bytes):
    return load_workbook(BytesIO(contenido)).active


class GenerarXlsxTests(SimpleTestCase):
    def test_primera_fila_son_los_encabezados_en_negrita_y_luego_los_datos(self):
        hoja = _abrir(generar_xlsx('Datos', ['Nombre', 'Total'], [['Ana', 3], ['Luis', 5]]))

        self.assertEqual([celda.value for celda in hoja[1]], ['Nombre', 'Total'])
        self.assertTrue(hoja['A1'].font.bold)
        self.assertEqual([[celda.value for celda in fila] for fila in hoja.iter_rows(min_row=2)], [['Ana', 3], ['Luis', 5]])

    def test_el_texto_que_empieza_como_formula_se_neutraliza(self):
        hoja = _abrir(generar_xlsx('Datos', ['Archivo'], [['=HYPERLINK("http://malo")'], ['+1+1'], ['-2'], ['@SUM(A1)'], ['normal']]))

        valores = [fila[0].value for fila in hoja.iter_rows(min_row=2)]

        self.assertEqual(valores[:4], ["'=HYPERLINK(\"http://malo\")", "'+1+1", "'-2", "'@SUM(A1)"])
        self.assertEqual(valores[4], 'normal')
        self.assertTrue(all(fila[0].data_type == 's' for fila in hoja.iter_rows(min_row=2)))

    def test_las_fechas_con_zona_horaria_se_guardan_como_fecha_con_formato(self):
        hoja = _abrir(generar_xlsx('Datos', ['Fecha'], [[datetime(2026, 10, 27, 20, 30, tzinfo=zona_utc.utc)]]))

        celda = hoja['A2']

        self.assertIsInstance(celda.value, datetime)
        self.assertEqual(celda.number_format, 'dd/mm/yyyy hh:mm')

    def test_los_valores_vacios_quedan_vacios(self):
        hoja = _abrir(generar_xlsx('Datos', ['A', 'B'], [['x', None]]))

        self.assertIsNone(hoja['B2'].value)

    def test_el_nombre_de_la_hoja_se_recorta_a_31_caracteres(self):
        hoja = _abrir(generar_xlsx('H' * 50, ['A'], []))

        self.assertEqual(len(hoja.title), 31)

    def test_deja_fija_la_fila_de_encabezados_y_activa_el_filtro(self):
        hoja = _abrir(generar_xlsx('Datos', ['A', 'B'], [['x', 'y']]))

        self.assertEqual(hoja.freeze_panes, 'A2')
        self.assertEqual(hoja.auto_filter.ref, 'A1:B2')

    def test_el_ancho_de_las_columnas_sigue_al_contenido_con_un_tope(self):
        hoja = _abrir(generar_xlsx('Datos', ['Corto', 'Largo'], [['a', 'x' * 200]]))

        self.assertLessEqual(hoja.column_dimensions['B'].width, 62)
        self.assertGreater(hoja.column_dimensions['B'].width, hoja.column_dimensions['A'].width)


class RespuestaXlsxTests(SimpleTestCase):
    def test_se_descarga_como_adjunto_con_el_tipo_de_excel(self):
        respuesta = respuesta_xlsx('informe.xlsx', b'datos')

        self.assertEqual(respuesta['Content-Type'], TIPO_XLSX)
        self.assertEqual(respuesta['Content-Disposition'], 'attachment; filename="informe.xlsx"')
        self.assertEqual(respuesta.content, b'datos')
