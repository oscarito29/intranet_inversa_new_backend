"""Archivos de Excel para las exportaciones. Solo arma el archivo; qué filas salen lo decide el selector de cada app."""

from collections.abc import Iterable, Sequence
from datetime import date, datetime
from io import BytesIO
from typing import Any

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

TIPO_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
ANCHO_MAXIMO_DE_COLUMNA = 60

# Excel ejecuta como fórmula cualquier celda de texto que empiece así. Los nombres de archivo y los motivos
# los escriben personas, así que se neutralizan con un apóstrofo (inyección de fórmulas / CSV injection).
_PREFIJOS_DE_FORMULA = ('=', '+', '-', '@', '\t', '\r')


def _valor_de_celda(valor: Any) -> Any:
    if isinstance(valor, datetime):
        # Excel no admite fechas con zona horaria: se guarda la hora local de quien la ve.
        return timezone.localtime(valor).replace(tzinfo=None) if timezone.is_aware(valor) else valor
    if isinstance(valor, str) and valor.startswith(_PREFIJOS_DE_FORMULA):
        return f"'{valor}"
    return valor


def generar_xlsx(hoja: str, encabezados: Sequence[str], filas: Iterable[Sequence[Any]]) -> bytes:
    libro = Workbook()
    pagina = libro.active
    if pagina is None:  # un libro nuevo siempre trae una hoja; esto solo le dice al verificador de tipos que existe
        pagina = libro.create_sheet()
    pagina.title = hoja[:31]  # Excel limita el nombre de la hoja a 31 caracteres

    pagina.append(list(encabezados))
    for celda in pagina[1]:
        celda.font = Font(bold=True)
        celda.alignment = Alignment(vertical='center')

    anchos = [len(encabezado) for encabezado in encabezados]
    for fila in filas:
        valores = [_valor_de_celda(valor) for valor in fila]
        pagina.append(valores)
        for posicion, valor in enumerate(valores):
            if isinstance(valor, datetime):
                pagina.cell(row=pagina.max_row, column=posicion + 1).number_format = 'dd/mm/yyyy hh:mm'
                anchos[posicion] = max(anchos[posicion], 16)
            elif isinstance(valor, date):
                pagina.cell(row=pagina.max_row, column=posicion + 1).number_format = 'dd/mm/yyyy'
                anchos[posicion] = max(anchos[posicion], 10)
            elif valor is not None:
                anchos[posicion] = max(anchos[posicion], len(str(valor)))

    for posicion, ancho in enumerate(anchos, start=1):
        pagina.column_dimensions[get_column_letter(posicion)].width = min(ancho + 2, ANCHO_MAXIMO_DE_COLUMNA)
    pagina.freeze_panes = 'A2'
    pagina.auto_filter.ref = pagina.dimensions

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def respuesta_xlsx(nombre_de_archivo: str, contenido: bytes) -> HttpResponse:
    respuesta = HttpResponse(contenido, content_type=TIPO_XLSX)
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre_de_archivo}"'
    return respuesta
