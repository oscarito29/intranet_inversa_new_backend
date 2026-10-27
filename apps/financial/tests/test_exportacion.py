from io import BytesIO

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.exportacion import TIPO_XLSX
from apps.core.testing import ApiTestCase, cliente_de, crear_usuario
from apps.financial.models import Documento

PERIODO = {'anio': 2026, 'mes': 10, 'corporacion': '001'}


def _documento(estado=Documento.PENDIENTE, financiera_id='001', **extra) -> Documento:
    return Documento.objects.create(
        nombre_archivo=extra.pop('nombre_archivo', 'estado.pdf'),
        ruta_archivo='/media/documentos/estado.pdf',
        doc_aprobado=estado,
        anio_id=2026,
        mes_id=10,
        corporacion_id='001',
        financiera_id=financiera_id,
        fecha=timezone.now(),
        email_usuario_sube='sube@inversa.hn',
        **extra,
    )


def _filas(respuesta) -> list[list]:
    hoja = load_workbook(BytesIO(respuesta.content)).active
    return [[celda.value for celda in fila] for fila in hoja.iter_rows(min_row=2)]


class ExportarPeriodoTests(ApiTestCase):
    def _exportar(self, cliente=None, **parametros):
        return (cliente or self.cliente_admin).get(reverse('financiero-exportar'), {**PERIODO, **parametros})

    def test_exige_autenticacion(self):
        self.assertEqual(self._exportar(cliente_de(None)).status_code, 401)

    def test_descarga_un_excel_con_nombre_que_dice_el_periodo(self):
        respuesta = self._exportar()

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], TIPO_XLSX)
        self.assertEqual(respuesta['Content-Disposition'], 'attachment; filename="estados-financieros-2026-10-corporacion-001.xlsx"')

    def test_el_encabezado_y_las_filas_son_lo_que_muestra_la_pantalla(self):
        _documento(estado=Documento.RECHAZADO, financiera_id='001', nombre_archivo='a.pdf', motivo_rechazo='Ilegible')
        _documento(estado=Documento.APROBADO, financiera_id='002', nombre_archivo='b.pdf')

        respuesta = self._exportar()
        hoja = load_workbook(BytesIO(respuesta.content)).active
        filas = _filas(respuesta)

        self.assertEqual([celda.value for celda in hoja[1]], ['Financiera', 'Estado', 'Archivo', 'Fecha de subida', 'Motivo'])
        self.assertEqual(
            [(fila[0], fila[1], fila[2], fila[4]) for fila in filas],
            [('FINANCIERA A', 'Rechazado', 'a.pdf', 'Ilegible'), ('FINANCIERA B', 'Aprobado', 'b.pdf', None)],
        )

    def test_una_financiera_sin_documento_sale_como_sin_cargar(self):
        filas = _filas(self._exportar())

        self.assertEqual([(fila[0], fila[1], fila[2]) for fila in filas], [('FINANCIERA A', 'Sin cargar', None), ('FINANCIERA B', 'Sin cargar', None)])

    def test_un_documento_sin_estado_cuenta_como_pendiente(self):
        _documento(estado=None)

        self.assertEqual(_filas(self._exportar())[0][1], 'Pendiente')

    def test_un_usuario_de_entidad_no_exporta(self):
        _documento(financiera_id='002', nombre_archivo='ajeno.pdf')

        self.assertEqual(self._exportar(self.cliente_usuario).status_code, 403)

    def test_aprobador_y_consulta_exportan_todas_las_entidades(self):
        _documento(financiera_id='001', nombre_archivo='a.pdf')
        _documento(financiera_id='002', nombre_archivo='b.pdf')

        for rol in (3, 4):
            with self.subTest(rol=rol):
                cliente = cliente_de(crear_usuario(rol_id=rol, email=f'rol{rol}@inversa.hn'))
                filas = _filas(self._exportar(cliente))
                self.assertEqual([(fila[0], fila[2]) for fila in filas], [('FINANCIERA A', 'a.pdf'), ('FINANCIERA B', 'b.pdf')])

    def test_el_motivo_de_una_revocacion_sale_en_la_columna_de_motivo(self):
        _documento(estado=Documento.REVOCADO, motivo_desaprueba='Cifras erradas')

        self.assertEqual(_filas(self._exportar())[0][4], 'Cifras erradas')

    def test_un_nombre_de_archivo_que_parece_formula_no_se_ejecuta(self):
        _documento(nombre_archivo='=CMD()')

        self.assertEqual(_filas(self._exportar())[0][2], "'=CMD()")

    def test_parametros_invalidos_dan_400(self):
        self.assertEqual(self._exportar(mes=13).status_code, 400)
        self.assertEqual(self._exportar(corporacion='999').status_code, 400)
