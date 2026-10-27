from io import BytesIO
from unittest import mock

from django.urls import reverse
from openpyxl import load_workbook

from apps.core.exportacion import TIPO_XLSX
from apps.core.testing import ApiTestCase, cliente_de
from apps.documentos import selectors
from apps.documentos.models import DocsEtiqueta, Etiqueta, EtiquetaRol

from .test_documentos import _categoria, _documento


def _filas(respuesta) -> list[list]:
    hoja = load_workbook(BytesIO(respuesta.content)).active
    return [[celda.value for celda in fila] for fila in hoja.iter_rows(min_row=2)]


class ExportarDocumentosTests(ApiTestCase):
    def setUp(self):
        super().setUp()
        self.contabilidad = _categoria('Contabilidad')
        self.rrhh = _categoria('RRHH')
        self.general = Etiqueta.objects.create(etiqueta_nombre='General')
        self.interna = Etiqueta.objects.create(etiqueta_nombre='Interna')
        EtiquetaRol.objects.create(etiqueta=self.general, rol_id=1)
        EtiquetaRol.objects.create(etiqueta=self.general, rol_id=2)
        EtiquetaRol.objects.create(etiqueta=self.interna, rol_id=1)

    def _exportar(self, cliente=None, **parametros):
        return (cliente or self.cliente_usuario).get(reverse('documentos-exportar'), parametros)

    def test_exige_autenticacion(self):
        self.assertEqual(self._exportar(cliente_de(None)).status_code, 401)

    def test_descarga_un_excel_con_los_encabezados_de_la_lista(self):
        _documento(self.contabilidad, nombre='balance.pdf')

        respuesta = self._exportar()
        hoja = load_workbook(BytesIO(respuesta.content)).active

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], TIPO_XLSX)
        self.assertEqual([celda.value for celda in hoja[1]], ['Documento', 'Categoría', 'Etiquetas', 'Fecha de subida', 'Tipo', 'Tamaño'])

    def test_cada_fila_trae_categoria_y_etiquetas_por_nombre(self):
        documento = _documento(self.contabilidad, nombre='balance.pdf')
        DocsEtiqueta.objects.create(documento=documento, etiqueta=self.interna)
        DocsEtiqueta.objects.create(documento=documento, etiqueta=self.general)

        fila = _filas(self._exportar(self.cliente_admin))[0]

        self.assertEqual((fila[0], fila[1], fila[2], fila[4], fila[5]), ('balance.pdf', 'Contabilidad', 'General, Interna', 'application/pdf', '1 KB'))

    def test_aplica_los_mismos_filtros_que_la_lista(self):
        _documento(self.contabilidad, nombre='balance.pdf')
        _documento(self.rrhh, nombre='nomina.pdf')

        por_categoria = _filas(self._exportar(categoria=self.rrhh.id))
        por_texto = _filas(self._exportar(q='balance'))

        self.assertEqual([fila[0] for fila in por_categoria], ['nomina.pdf'])
        self.assertEqual([fila[0] for fila in por_texto], ['balance.pdf'])

    def test_exporta_todos_y_no_una_pagina(self):
        for numero in range(30):
            _documento(self.contabilidad, nombre=f'doc-{numero:02d}.pdf')

        self.assertEqual(len(_filas(self._exportar(page_size=5))), 30)

    def test_no_incluye_los_documentos_reemplazados(self):
        _documento(self.contabilidad, nombre='vigente.pdf')
        _documento(self.contabilidad, estado='Reemplazado', nombre='viejo.pdf')

        self.assertEqual([fila[0] for fila in _filas(self._exportar())], ['vigente.pdf'])

    def test_un_usuario_no_exporta_lo_que_su_rol_no_puede_ver(self):
        publico = _documento(self.contabilidad, nombre='publico.pdf')
        reservado = _documento(self.contabilidad, nombre='reservado.pdf')
        DocsEtiqueta.objects.create(documento=publico, etiqueta=self.general)
        DocsEtiqueta.objects.create(documento=reservado, etiqueta=self.interna)

        usuario = _filas(self._exportar(self.cliente_usuario))
        admin = _filas(self._exportar(self.cliente_admin))

        self.assertEqual([fila[0] for fila in usuario], ['publico.pdf'])
        self.assertEqual({fila[0] for fila in admin}, {'publico.pdf', 'reservado.pdf'})

    def test_orden_por_fecha_como_la_lista(self):
        _documento(self.contabilidad, nombre='primero.pdf')
        _documento(self.contabilidad, nombre='segundo.pdf')

        recientes = [fila[0] for fila in _filas(self._exportar(orden='recientes'))]
        antiguos = [fila[0] for fila in _filas(self._exportar(orden='antiguos'))]

        self.assertEqual(recientes, ['segundo.pdf', 'primero.pdf'])
        self.assertEqual(antiguos, ['primero.pdf', 'segundo.pdf'])

    def test_con_demasiados_documentos_pide_acotar_el_filtro(self):
        for numero in range(3):
            _documento(self.contabilidad, nombre=f'doc-{numero}.pdf')

        with mock.patch.object(selectors, 'EXPORTACION_MAXIMA', 2):
            respuesta = self._exportar()

        self.assertEqual(respuesta.status_code, 400)
        self.assertIn('Acótalos', respuesta.data['detail'])

    def test_un_nombre_de_documento_que_parece_formula_no_se_ejecuta(self):
        _documento(self.contabilidad, nombre='=SUMA(1;1).pdf')

        self.assertEqual(_filas(self._exportar())[0][0], "'=SUMA(1;1).pdf")

    def test_etiquetas_invalidas_dan_400(self):
        self.assertEqual(self._exportar(etiquetas='a,b').status_code, 400)
